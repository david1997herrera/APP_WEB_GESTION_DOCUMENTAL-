"""Libro de Excel con el banner de Rosemirovich Roses."""

import os
from datetime import datetime
from io import BytesIO

from flask import current_app
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ImagenHoja
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AZUL = "0010A0"
ROJO = "C00000"
FILA_ENCABEZADO = 6


def construir_libro(titulo, columnas, filas):
    """Arma un libro con logo, franja de marca y la tabla pedida."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = (titulo or "Reporte")[:31]
    cantidad = max(len(columnas), 4)
    ultima = get_column_letter(cantidad)

    _banner(hoja, titulo, ultima, cantidad)
    _encabezados(hoja, columnas)
    _filas(hoja, filas, cantidad)
    _anchos(hoja, columnas, filas)

    hoja.freeze_panes = "A7"
    hoja.auto_filter.ref = f"A{FILA_ENCABEZADO}:{ultima}{FILA_ENCABEZADO + max(len(filas), 1)}"
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.page_setup.paperSize = hoja.PAPERSIZE_A4
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.oddHeader.left.text = "Rosemirovich Roses"
    hoja.oddFooter.center.text = "Roses around the world · Gestión Documental"
    hoja.print_title_rows = "1:6"
    hoja.sheet_view.showGridLines = False

    salida = BytesIO()
    libro.save(salida)
    salida.seek(0)
    return salida


def _banner(hoja, titulo, ultima, cantidad):
    azul = PatternFill("solid", fgColor=AZUL)
    rojo = PatternFill("solid", fgColor=ROJO)
    blanco = Font(name="Calibri", color="FFFFFF", bold=True, size=16)
    subtitulo = Font(name="Calibri", color="FFFFFF", size=11)
    detalle = Font(name="Calibri", color="FFFFFF", size=10)

    for fila in range(1, 5):
        for columna in range(1, cantidad + 1):
            hoja.cell(fila, columna).fill = azul
        hoja.row_dimensions[fila].height = 18
    hoja.row_dimensions[1].height = 22
    hoja.row_dimensions[2].height = 20

    hoja.merge_cells(start_row=1, start_column=3, end_row=1, end_column=cantidad)
    hoja.merge_cells(start_row=2, start_column=3, end_row=2, end_column=cantidad)
    hoja.merge_cells(start_row=3, start_column=3, end_row=3, end_column=cantidad)
    nombre = hoja.cell(1, 3, "ROSEMIROVICH ROSES")
    nombre.font = blanco
    nombre.alignment = Alignment(vertical="center")
    frase = hoja.cell(2, 3, "Roses around the world")
    frase.font = subtitulo
    frase.alignment = Alignment(vertical="center")
    reporte = hoja.cell(
        3,
        3,
        f"{titulo} · {datetime.now().strftime('%d/%m/%Y %H:%M')}",
    )
    reporte.font = detalle
    reporte.alignment = Alignment(vertical="center")

    for columna in range(1, cantidad + 1):
        hoja.cell(5, columna).fill = rojo
    hoja.row_dimensions[5].height = 6

    ruta = os.path.join(current_app.static_folder, "img", "logo.png")
    if os.path.isfile(ruta):
        imagen = ImagenHoja(ruta)
        imagen.width = 118
        imagen.height = 80
        hoja.add_image(imagen, "A1")


def _encabezados(hoja, columnas):
    relleno = PatternFill("solid", fgColor=AZUL)
    letra = Font(name="Calibri", color="FFFFFF", bold=True, size=11)
    borde = Border(
        bottom=Side(style="thin", color=ROJO),
    )
    hoja.row_dimensions[FILA_ENCABEZADO].height = 22
    for indice, titulo in enumerate(columnas, start=1):
        celda = hoja.cell(FILA_ENCABEZADO, indice, titulo)
        celda.fill = relleno
        celda.font = letra
        celda.alignment = Alignment(vertical="center", horizontal="left")
        celda.border = borde


def _filas(hoja, filas, cantidad):
    cebra = PatternFill("solid", fgColor="F4F6FB")
    letra = Font(name="Calibri", size=11, color="1C1C1C")
    for desplazamiento, valores in enumerate(filas):
        numero = FILA_ENCABEZADO + 1 + desplazamiento
        for indice in range(1, cantidad + 1):
            valor = valores[indice - 1] if indice - 1 < len(valores) else ""
            celda = hoja.cell(numero, indice, valor if valor is not None else "")
            celda.font = letra
            celda.alignment = Alignment(vertical="center")
            if desplazamiento % 2 == 1:
                celda.fill = cebra


def _anchos(hoja, columnas, filas):
    for indice, titulo in enumerate(columnas, start=1):
        mayor = len(str(titulo))
        for fila in filas[:200]:
            if indice - 1 < len(fila) and fila[indice - 1] is not None:
                mayor = max(mayor, min(len(str(fila[indice - 1])), 48))
        hoja.column_dimensions[get_column_letter(indice)].width = max(14, min(mayor + 3, 42))
    hoja.column_dimensions["A"].width = 18
    hoja.column_dimensions["B"].width = 16
